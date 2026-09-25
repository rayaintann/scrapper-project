"""Import workbook full_classification_fill_100.xlsx (100 KOL x 10 field) ke DB `kol` -- hanya ke storage LABEL yang sah.

    python classification_workbook_import.py <workbook.xlsx>                     DRY-RUN (default, read-only)
    python classification_workbook_import.py <workbook.xlsx> --apply             satu transaksi; validasi; COMMIT/ROLLBACK
    python classification_workbook_import.py <workbook.xlsx> --only "Creator Gender" [--apply]

Pemetaan field -> storage label (bukan cacah/metric):

    Creator Style / Personality   public.kol_attribute_map, source='curated', confidence, evidence jsonb.
                                  Writer classifier tidak menulis kind yang punya baris curated.
    Creator Gender                l2_gold.kol_profile_card.creator_gender, source='manual' (kolom label; asset
                                  creator_gender tidak pernah menimpa 'manual').
    Creator Category/Subcategory  public.kol_directory.inferred_*, source='curated' (migration 054). Writer classifier
                                  melewati baris 'curated'. Category existing dipertahankan; evidence classifier lama
                                  disimpan di inferred_category_evidence.preserved.
    Audience Gender/Age/Country/City
                                  feature.{ig,tt}_audience_analysis.curated_* (migration 054): kolom LABEL di baris
                                  Analysis Audience existing, terpisah dari kolom measurement; asset audience_feature
                                  tidak menulisnya. audience_classification memakainya hanya bila measured/inferred Unknown.
    Audience Interest             sudah 100/100 dari Audience Analysis existing -> SKIP_SAME.

Aturan: nilai existing dipertahankan (SKIP_SAME); hanya gap/NULL yang diisi; tidak ada DELETE; import ulang = 0 write.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FIELDS = ("Audience Gender", "Audience Interest", "Creator Gender", "Creator Category", "Creator Subcategory",
          "Creator Style", "Creator Personality", "Audience Age", "Audience Country", "Audience City")
WRITABLE = {"Creator Style": "style", "Creator Personality": "personality"}
CARD_FIELDS = {"Creator Gender": "creator_gender"}
KD_FIELDS = ("Creator Category", "Creator Subcategory")
AUD_FIELDS = {"Audience Gender": "gender", "Audience Age": "age", "Audience Country": "country",
              "Audience City": "city"}
READ_ONLY_FIELDS = {"Audience Interest": "interest"}          # sudah lengkap; tidak ada storage kurasi
#: Dipertahankan untuk kompatibilitas: tidak ada lagi field yang diblokir storage-nya.
BLOCKED: dict[str, str] = {}
GENDERS = ("female", "male")
AUD_GENDERS = ("female", "male", "balanced")
AGE_BUCKETS = ("13-17", "18-24", "25-34", "35-44", "45+")
CONF = ("high", "medium", "low")
CURATED = "curated"


@dataclass
class Cell:
    kol_id: str
    username: str
    platform: str
    field: str
    final: str
    method: str
    confidence: str = ""
    evidence: str = ""
    source: str = ""
    reason: str = ""
    social_account_id: str = ""


@dataclass
class Item:
    cell: Cell
    status: str          # SKIP_SAME | INSERT | UPDATE | NOT_SELECTED | INVALID
    reason: str = ""
    params: dict = field(default_factory=dict)


@dataclass
class DbState:
    """Snapshot DB yang dibutuhkan plan (murni setelah dimuat)."""
    kol_ids: set[str]
    attributes: dict[tuple[str, str], str]                    # (kind, key) -> kol_attribute.id
    attr_map: dict[tuple[str, str], set[str]]                 # (kol_id, kind) -> {attribute_key}
    card: dict[str, tuple] = field(default_factory=dict)      # sa -> (creator_gender, source)
    kd: dict[str, dict] = field(default_factory=dict)         # kol_id -> inferred_* state
    categories: dict[str, str] = field(default_factory=dict)  # category name -> id (canonical)
    subcats: dict[tuple[str, str], str] = field(default_factory=dict)   # (category id, sub name) -> id
    cat_names: dict[str, str] = field(default_factory=dict)   # id -> name (category + sub)
    aud_curated: dict[str, dict] = field(default_factory=dict)  # sa -> {attr: value} dari feature.curated_*
    aud_state: dict[str, dict] = field(default_factory=dict)  # sa -> {attr: value dari audience_classification}
    cities: set[str] = field(default_factory=set)
    has_curated_cols: bool = True                              # migration 054 sudah diterapkan?


def read_workbook(path: Path) -> list[Cell]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    rows = list(wb["FULL_CLASSIFICATION"].iter_rows(values_only=True))
    h = list(rows[0])
    ev_rows = list(wb["EVIDENCE"].iter_rows(values_only=True))
    eh = list(ev_rows[0])
    ev = {}
    for r in ev_rows[1:]:
        d = dict(zip(eh, r))
        ev[(d["kol_id"], d["field"])] = d
    out = []
    for r in rows[1:]:
        for f in FIELDS:
            e = ev.get((r[h.index("kol_id")], f), {})
            out.append(Cell(r[h.index("kol_id")], r[h.index("username")], r[h.index("platform")], f,
                            r[h.index(f)], r[h.index(f + " [method]")], e.get("confidence") or "",
                            e.get("evidence") or "", e.get("source") or "", e.get("reason") or "",
                            r[h.index("social_account_id")]))
    return out


def _conf(c: Cell) -> str:
    return c.confidence if c.confidence in CONF else "low"


def _ev(c: Cell) -> dict:
    return {"method": c.method, "confidence": _conf(c), "evidence": (c.evidence or "")[:400],
            "source": c.source, "reason": (c.reason or "")[:300]}


# --- plan per jalur (murni) -------------------------------------------------------
def _plan_gender(c: Cell, db: DbState) -> Item:
    if c.social_account_id not in db.card:
        return Item(c, "INVALID", "akun tidak punya baris kol_profile_card")
    have, src = db.card[c.social_account_id]
    if c.final not in GENDERS:
        return Item(c, "INVALID", f"gender harus salah satu {GENDERS}")
    if have == c.final:
        return Item(c, "SKIP_SAME", "sudah ada di DB")
    if have is not None or c.method == "EXISTING":
        return Item(c, "INVALID", f"DB={have} ({src}); nilai existing tidak ditimpa")
    return Item(c, "UPDATE", "", {"sa": c.social_account_id, "gender": c.final, "confidence": _conf(c)})


def _plan_attr(c: Cell, db: DbState) -> Item:
    kind = WRITABLE[c.field]
    have = db.attr_map.get((c.kol_id, kind), set())
    if c.final in have:
        return Item(c, "SKIP_SAME", "sudah ada di DB")
    if c.method == "EXISTING":
        return Item(c, "INVALID", f"workbook EXISTING={c.final} tapi DB={sorted(have)}")
    if have:
        return Item(c, "INVALID", f"DB sudah berisi {sorted(have)}; nilai existing tidak ditimpa")
    aid = db.attributes.get((kind, c.final))
    if aid is None:
        return Item(c, "INVALID", f"{c.final} bukan kol_attribute aktif kind '{kind}'")
    return Item(c, "INSERT", "", {
        "kol": c.kol_id, "attr": aid, "confidence": _conf(c),
        "note": "full_classification_fill_100.xlsx (direview user 2026-09-25)",
        "evidence": {"basis": "workbook_review", **_ev(c)}})


def _plan_kd(c: Cell, db: DbState, cat_final: dict[str, str]) -> Item:
    st = db.kd.get(c.kol_id)
    if st is None:
        return Item(c, "INVALID", "kol_id tidak ada di kol_directory")
    if c.field == "Creator Category":
        have = db.cat_names.get(st["cat"]) if st["cat"] else None
        if have == c.final:
            return Item(c, "SKIP_SAME", "sudah ada di DB")
        if have is not None or c.method == "EXISTING":
            return Item(c, "INVALID", f"DB category={have}; nilai existing tidak ditimpa")
        cid = db.categories.get(c.final)
        if cid is None:
            return Item(c, "INVALID", f"{c.final} bukan kategori canonical")
        return Item(c, "UPDATE", "", {"cat": cid, "cat_conf": _conf(c), "ev": _ev(c)})
    have = db.cat_names.get(st["sub"]) if st["sub"] else None
    if have == c.final:
        return Item(c, "SKIP_SAME", "sudah ada di DB")
    if have is not None or c.method == "EXISTING":
        return Item(c, "INVALID", f"DB subcategory={have}; nilai existing tidak ditimpa")
    parent = st["cat"] or db.categories.get(cat_final.get(c.kol_id, ""))
    sid = db.subcats.get((parent, c.final)) if parent else None
    if sid is None:
        return Item(c, "INVALID", f"{c.final} bukan subkategori dari {cat_final.get(c.kol_id)}")
    return Item(c, "UPDATE", "", {"sub": sid, "sub_conf": _conf(c), "ev": _ev(c)})


def _plan_audience(c: Cell, db: DbState) -> Item:
    attr = AUD_FIELDS[c.field]
    measured = db.aud_state.get(c.social_account_id, {}).get(attr)
    have = db.aud_curated.get(c.social_account_id, {}).get(attr)
    if measured == c.final or have == c.final:
        return Item(c, "SKIP_SAME", "sudah ada di DB")
    if c.method == "EXISTING" or measured is not None:
        return Item(c, "INVALID", f"DB={measured}; nilai existing tidak ditimpa")
    if have is not None:
        return Item(c, "INVALID", f"label kurasi DB={have}; tidak ditimpa")
    ok = {"gender": c.final in AUD_GENDERS, "age": c.final in AGE_BUCKETS,
          "country": isinstance(c.final, str) and len(c.final) == 2 and c.final.isalpha() and c.final.isupper(),
          "city": c.final in db.cities}[attr]
    if not ok:
        return Item(c, "INVALID", f"{c.final} di luar taxonomy {attr}")
    return Item(c, "UPDATE", "", {"sa": c.social_account_id, "platform": c.platform, "attr": attr,
                                  "value": c.final, "ev": _ev(c)})


def plan(cells: list[Cell], db: DbState, only: set[str] | None = None) -> list[Item]:
    items, seen = [], set()
    cat_final = {c.kol_id: c.final for c in cells if c.field == "Creator Category"}
    for c in cells:
        if c.kol_id not in db.kol_ids:
            items.append(Item(c, "INVALID", "kol_id tidak ada di kol_directory")); continue
        if (c.kol_id, c.field) in seen:
            items.append(Item(c, "INVALID", "duplikat (kol_id, field) di workbook")); continue
        seen.add((c.kol_id, c.field))
        if only is not None and c.field not in only:
            items.append(Item(c, "NOT_SELECTED", "di luar --only")); continue
        if c.field in CARD_FIELDS:
            items.append(_plan_gender(c, db))
        elif c.field in WRITABLE:
            items.append(_plan_attr(c, db))
        elif c.field in KD_FIELDS:
            items.append(_plan_kd(c, db, cat_final))
        elif c.field in AUD_FIELDS:
            items.append(_plan_audience(c, db))
        else:                                                    # Audience Interest
            have = db.aud_state.get(c.social_account_id, {}).get(READ_ONLY_FIELDS[c.field])
            items.append(Item(c, "SKIP_SAME", "") if have == c.final else
                         Item(c, "INVALID", f"interest DB={have}; tidak ada storage kurasi"))
    return items


def kd_writes(items: list[Item], db: DbState) -> list[dict]:
    """Gabungkan item category/subcategory per KOL jadi satu UPDATE kol_directory."""
    per: dict[str, dict] = defaultdict(dict)
    for it in items:
        if it.status == "UPDATE" and it.cell.field in KD_FIELDS:
            per[it.cell.kol_id][it.cell.field] = it
    out = []
    for kid, f in per.items():
        st = db.kd[kid]
        cat_it, sub_it = f.get("Creator Category"), f.get("Creator Subcategory")
        cat = cat_it.params["cat"] if cat_it else st["cat"]
        out.append({
            "kid": kid, "expect_cat": st["cat"], "cat": cat,
            "sub": sub_it.params["sub"] if sub_it else st["sub"],
            "cat_conf": cat_it.params["cat_conf"] if cat_it else st["cat_conf"],
            "sub_conf": sub_it.params["sub_conf"] if sub_it else st["sub_conf"],
            "evidence": {"basis": "workbook_review", "source": CURATED,
                         "category": cat_it.params["ev"] if cat_it else {"method": "EXISTING", "preserved": True},
                         "subcategory": sub_it.params["ev"] if sub_it else None,
                         "preserved": st["evidence"]}})
    return out


def aud_writes(items: list[Item]) -> list[dict]:
    per: dict[str, dict] = defaultdict(lambda: {"gender": None, "age": None, "country": None, "city": None,
                                                "evidence": {}, "table": None})
    for it in items:
        if it.status == "UPDATE" and it.cell.field in AUD_FIELDS:
            row = per[it.params["sa"]]
            row["table"] = AUD_TABLE[it.params["platform"]]
            row[it.params["attr"]] = it.params["value"]
            row["evidence"][it.params["attr"]] = it.params["ev"]
    return [{"sa": sa, **v} for sa, v in per.items()]


# --- SQL ------------------------------------------------------------------------
SQL_KOL = "SELECT id::text FROM public.kol_directory WHERE id = ANY(%s::uuid[])"
SQL_ATTR = "SELECT kind, attribute_key, id::text FROM public.kol_attribute WHERE is_active"
SQL_MAP = """SELECT m.kol_directory_id::text, a.kind, a.attribute_key FROM public.kol_attribute_map m
               JOIN public.kol_attribute a ON a.id = m.kol_attribute_id WHERE m.kol_directory_id = ANY(%s::uuid[])"""
SQL_INSERT = """
    INSERT INTO public.kol_attribute_map (kol_directory_id, kol_attribute_id, source, confidence, note, evidence)
    SELECT %(kol)s, %(attr)s, 'curated', %(confidence)s, %(note)s, %(evidence)s
     WHERE NOT EXISTS (SELECT 1 FROM public.kol_attribute_map
                        WHERE kol_directory_id = %(kol)s AND kol_attribute_id = %(attr)s)"""
SQL_CARD = """SELECT social_account_id::text, creator_gender, creator_gender_source FROM l2_gold.kol_profile_card
               WHERE social_account_id = ANY(%s::uuid[])"""
SQL_GENDER = """
    UPDATE l2_gold.kol_profile_card
       SET creator_gender = %(gender)s, creator_gender_source = 'manual', creator_gender_confidence = %(confidence)s
     WHERE social_account_id = %(sa)s AND creator_gender IS NULL
       AND creator_gender_source IS DISTINCT FROM 'manual'"""
SQL_KD_STATE = """SELECT id::text, inferred_category_id::text, inferred_subcategory_id::text, inferred_category_source,
                         inferred_category_confidence, inferred_subcategory_confidence, inferred_category_evidence
                    FROM public.kol_directory WHERE id = ANY(%s::uuid[])"""
SQL_CATS = """SELECT id::text, name, level, parent_id::text FROM public.kol_categories"""
SQL_KD_CURATED = """
    UPDATE public.kol_directory
       SET inferred_category_id = %(cat)s, inferred_subcategory_id = %(sub)s,
           inferred_category_source = 'curated',
           inferred_category_confidence = %(cat_conf)s, inferred_subcategory_confidence = %(sub_conf)s,
           inferred_category_evidence = %(evidence)s, inferred_category_at = now()
     WHERE id = %(kid)s
       AND inferred_category_id IS NOT DISTINCT FROM %(expect_cat)s::uuid
       AND inferred_subcategory_id IS NULL"""
AUD_TABLE = {"instagram": "feature.ig_audience_analysis", "tiktok": "feature.tt_audience_analysis"}
SQL_AUD_CURATED_STATE = """
    SELECT social_account_id::text, curated_gender, curated_age, curated_country, curated_city
      FROM feature.ig_audience_analysis WHERE social_account_id = ANY(%(sa)s::uuid[])
    UNION ALL
    SELECT social_account_id::text, curated_gender, curated_age, curated_country, curated_city
      FROM feature.tt_audience_analysis WHERE social_account_id = ANY(%(sa)s::uuid[])"""
#: Hanya kolom curated_* yang disentuh; kolom measurement dan updated_at TIDAK. COALESCE = hanya isi NULL.
SQL_AUD_UPDATE = """
    UPDATE {table}
       SET curated_gender   = COALESCE(curated_gender,  %(gender)s),
           curated_age      = COALESCE(curated_age,     %(age)s),
           curated_country  = COALESCE(curated_country, %(country)s),
           curated_city     = COALESCE(curated_city,    %(city)s),
           curated_evidence = %(evidence)s::jsonb || coalesce(curated_evidence, '{{}}'::jsonb)
     WHERE social_account_id = %(sa)s"""
#: Sidik isi yang TIDAK boleh berubah oleh import (kolom target dikecualikan dari tabelnya sendiri).
SQL_FINGERPRINT = {
    "kol_profile_card (tanpa creator_gender*)": """SELECT md5(string_agg((to_jsonb(c) - 'creator_gender'
        - 'creator_gender_source' - 'creator_gender_confidence')::text, '|' ORDER BY c.id::text))
        FROM l2_gold.kol_profile_card c""",
    "kol_directory (tanpa inferred_*)": """SELECT md5(string_agg((to_jsonb(k) - 'inferred_category_id'
        - 'inferred_subcategory_id' - 'inferred_category_source' - 'inferred_category_confidence'
        - 'inferred_subcategory_confidence' - 'inferred_category_evidence' - 'inferred_category_at')::text,
        '|' ORDER BY k.id::text)) FROM public.kol_directory k""",
    "kol_attribute_map": "SELECT md5(string_agg(to_jsonb(m)::text, '|' ORDER BY m.id::text)) FROM public.kol_attribute_map m",
    **{f"{t} (tanpa curated_*)": f"""SELECT md5(string_agg((to_jsonb(t) - 'curated_gender' - 'curated_age'
        - 'curated_country' - 'curated_city' - 'curated_evidence')::text, '|' ORDER BY t.id::text)) FROM {t} t"""
       for t in ("feature.ig_audience_analysis", "feature.tt_audience_analysis")},
}
HIST_SCHEMAS = ("l0_raw", "l0_harmonization", "l1_silver", "feature", "l2_gold")


def load_db(conn, cells: list[Cell]) -> DbState:
    import audience_classification as A
    import audience_inference as AI
    kol_ids = sorted({c.kol_id for c in cells})
    sas = sorted({c.social_account_id for c in cells})
    cur = conn.cursor()
    cur.execute(SQL_KOL, (kol_ids,)); found = {r[0] for r in cur.fetchall()}
    cur.execute(SQL_ATTR); attrs = {(k, key): i for k, key, i in cur.fetchall()}
    cur.execute(SQL_MAP, (kol_ids,))
    amap: dict[tuple[str, str], set[str]] = {}
    for kid, kind, key in cur.fetchall():
        amap.setdefault((kid, kind), set()).add(key)
    cur.execute(SQL_CARD, (sas,)); card = {r[0]: (r[1], r[2]) for r in cur.fetchall()}
    cur.execute(SQL_KD_STATE, (kol_ids,))
    kd = {r[0]: {"cat": r[1], "sub": r[2], "source": r[3], "cat_conf": r[4], "sub_conf": r[5], "evidence": r[6]}
          for r in cur.fetchall()}
    cur.execute(SQL_CATS)
    cats, subs, names = {}, {}, {}
    import creator_classification_lexicon as L
    rows = cur.fetchall()
    cur.execute("SELECT id::text, code FROM public.kol_categories")
    code = dict(cur.fetchall())
    for cid, name, level, parent in rows:
        names[cid] = name
        if level == "category" and parent is None and code.get(cid) in L.CATEGORY_TARGETS:
            cats[name] = cid
        if level == "sub_category" and parent:
            subs[(parent, name)] = cid
    aud_cur: dict[str, dict] = {}
    cur.execute(A.SQL_HAS_CURATED)
    has_cols = bool(cur.fetchone()[0])
    if has_cols:
        cur.execute(SQL_AUD_CURATED_STATE, {"sa": sas})
        for sa, g, age, co, ci in cur.fetchall():
            aud_cur[sa] = {"gender": g, "age": age, "country": co, "city": ci}
    # audience state = hasil measured/inferred SAJA (tanpa label kurasi) -> nilai "existing"
    ainp = A.load(conn)
    aud_state = {}
    for sa in sas:
        inp = ainp.get(sa, A.AudienceInput())
        raw = {"gender": A.classify_gender(inp), "age": A.classify_age(inp), "country": A.classify_country(inp),
               "city": A.classify_city(inp), "interest": A.classify_interest(inp)}
        aud_state[sa] = {k: v.value for k, v in raw.items()}
    cur.execute("SELECT DISTINCT geo_key FROM l2_gold.audience_geo_daily WHERE geo_level = 'city' AND geo_key <> 'unknown'")
    cities = set(AI.KOTA_ID.values()) | set(AI.PROVINSI_ID.values()) | {r[0] for r in cur.fetchall()}
    return DbState(found, attrs, amap, card, kd, cats, subs, names, aud_cur, aud_state, cities, has_cols)


def hist_counts(cur) -> dict[str, int]:
    cur.execute("SELECT schemaname || '.' || relname FROM pg_stat_user_tables WHERE schemaname = ANY(%s)",
                (list(HIST_SCHEMAS),))
    out = {}
    for (t,) in cur.fetchall():
        cur.execute(f"SELECT count(*) FROM {t}"); out[t] = cur.fetchone()[0]
    return out


def fingerprints(cur) -> dict[str, str]:
    out = {}
    for k, sql in SQL_FINGERPRINT.items():
        cur.execute(sql); out[k] = cur.fetchone()[0]
    return out


def summary(items: list[Item]) -> dict:
    per = {f: Counter() for f in FIELDS}
    for it in items:
        per[it.cell.field][it.status] += 1
    return {f: dict(c) for f, c in per.items()}


def coverage(conn) -> dict[str, int]:
    """Coverage 10 field dengan logic coverage existing (one_pass.read_state + audience_classification)."""
    sys.path.insert(0, str(ROOT / "orchestration"))
    import audience_classification as A
    from kol_orchestration import one_pass as OP
    members = OP.load_cohort()
    st = OP.read_state(conn, members, A.load(conn))
    keys = {"Audience Gender": ("audience", "gender"), "Audience Interest": ("audience", "interest"),
            "Creator Gender": ("creator", "gender"), "Creator Category": ("creator", "category"),
            "Creator Subcategory": ("creator", "subcategory"), "Creator Style": ("creator", "style"),
            "Creator Personality": ("creator", "personality"), "Audience Age": ("audience", "age"),
            "Audience Country": ("audience", "country"), "Audience City": ("audience", "city")}
    return {f: sum(1 for m in members if st[m.social_account_id][g][k] not in (None, "", "unknown", "Unknown"))
            for f, (g, k) in keys.items()}


def dry_run_report(cells: list[Cell], db: DbState, only: set[str] | None) -> tuple[list[Item], dict]:
    items = plan(cells, db, only)
    st = Counter(it.status for it in items)
    ins_keys = [(it.params["kol"], it.params["attr"]) for it in items
                if it.status == "INSERT" and it.cell.field in WRITABLE]
    upd_sa = [it.params["sa"] for it in items if it.status == "UPDATE" and it.cell.field in CARD_FIELDS]
    dup_target = (len(ins_keys) - len(set(ins_keys))) + (len(upd_sa) - len(set(upd_sa)))
    report = {"workbook_cells": len(cells), "unique_kol": len({c.kol_id for c in cells}),
              "fields": len({c.field for c in cells}), "only": sorted(only) if only else "semua field",
              "existing_values": sum(1 for c in cells if c.method == "EXISTING"),
              "insert": st["INSERT"], "update": st["UPDATE"], "skip_same": st["SKIP_SAME"],
              "not_selected": st["NOT_SELECTED"], "invalid": st["INVALID"], "duplicate_target": dup_target,
              "delete": 0, "kol_directory_rows": len(kd_writes(items, db)),
              "feature_audience_rows": len(aud_writes(items)),
              "migration_054_applied": db.has_curated_cols, "per_field": summary(items)}
    return items, report


def passes(report: dict) -> bool:
    return (not report["invalid"] and not report["duplicate_target"] and report["unique_kol"] == 100
            and report["workbook_cells"] == 1000)


def execute(conn, cells: list[Cell], only: set[str] | None) -> tuple[bool, dict, dict]:
    """Tulis di transaksi `conn` yang SEDANG berjalan lalu validasi. TIDAK commit/rollback --
    pemanggil yang memutuskan. Mengembalikan (lolos, checks, counts)."""
    import psycopg2.extras
    cur = conn.cursor()
    db = load_db(conn, cells)
    items, report = dry_run_report(cells, db, only)
    if not passes(report):
        return False, {"dry-run PASS": False}, {}
    kdw, audw = kd_writes(items, db), aud_writes(items)
    if audw and not db.has_curated_cols:
        return False, {"migration 054 diterapkan (kolom curated_* ada)": False}, {}
    n_attr_rows = sum(1 for it in items if it.status == "INSERT" and it.cell.field in WRITABLE)
    n_gender = sum(1 for it in items if it.status == "UPDATE" and it.cell.field in CARD_FIELDS)
    hist0, fp0 = hist_counts(cur), fingerprints(cur)
    stat = ("SELECT coalesce(sum(n_tup_del),0), coalesce(sum(n_tup_upd),0), coalesce(sum(n_tup_ins),0) "
            "FROM pg_stat_xact_user_tables")
    cur.execute(stat); d0, u0, i0 = cur.fetchone()
    n = Counter()
    for it in items:
        if it.status == "INSERT" and it.cell.field in WRITABLE:
            cur.execute(SQL_INSERT, {**it.params, "evidence": psycopg2.extras.Json(it.params["evidence"])})
            n["attr"] += cur.rowcount
        elif it.status == "UPDATE" and it.cell.field in CARD_FIELDS:
            cur.execute(SQL_GENDER, it.params); n["gender"] += cur.rowcount
    for w in kdw:
        cur.execute(SQL_KD_CURATED, {**w, "evidence": psycopg2.extras.Json(w["evidence"])}); n["kd"] += cur.rowcount
    for w in audw:
        cur.execute(SQL_AUD_UPDATE.format(table=w["table"]),
                    {**w, "evidence": psycopg2.extras.Json(w["evidence"])}); n["aud"] += cur.rowcount
    cur.execute(stat); d1, u1, i1 = cur.fetchone()
    fp1 = fingerprints(cur)
    cur.execute("""SELECT count(*) FROM (SELECT kol_directory_id, kol_attribute_id FROM public.kol_attribute_map
                    GROUP BY 1, 2 HAVING count(*) > 1) d""")
    dups = cur.fetchone()[0]
    keep = [k for k in fp0 if k != "kol_attribute_map" or n_attr_rows == 0]
    checks = {
        "attr inserted == plan": n["attr"] == n_attr_rows,
        "gender updated == plan": n["gender"] == n_gender,
        "kol_directory rows == plan": n["kd"] == len(kdw),
        "feature audience rows == plan": n["aud"] == len(audw),
        "no delete": d1 == d0,
        "row updates == plan": u1 - u0 == n_gender + len(kdw) + len(audw),
        "row inserts == plan": i1 - i0 == n_attr_rows,
        "no duplicate (kol, attribute)": dups == 0,
        "historical layer row counts unchanged": hist_counts(cur) == hist0,
        **{f"unchanged: {k}": fp1[k] == fp0[k] for k in keep},
    }
    if only is None:
        cov = coverage(conn)
        checks["10 field = 100/100"] = all(v == 100 for v in cov.values())
        n["coverage"] = cov
    return all(checks.values()), checks, dict(n)


def main(argv=None) -> int:
    import os
    import psycopg2
    from dotenv import load_dotenv
    p = argparse.ArgumentParser()
    p.add_argument("workbook", type=Path)
    p.add_argument("--apply", action="store_true")
    p.add_argument("--only", action="append", choices=FIELDS, default=None,
                   help="Proses hanya field ini (boleh diulang). Field lain -> NOT_SELECTED, tidak disentuh.")
    a = p.parse_args(argv)
    only = set(a.only) if a.only else None
    cells = read_workbook(a.workbook)
    load_dotenv(ROOT / ".env")
    conn = psycopg2.connect(host=os.environ["PG_HOST"], port=os.environ["PG_PORT"], dbname=os.environ["PG_DB"],
                            user=os.environ["PG_USER"], password=os.environ["PG_PASSWORD"], connect_timeout=15)
    conn.set_session(readonly=not a.apply)
    try:
        db = load_db(conn, cells)
        items, report = dry_run_report(cells, db, only)
        print(json.dumps(report, indent=1, ensure_ascii=False))
        for it in items:
            if it.status == "INVALID":
                print("INVALID", it.cell.username, it.cell.field, it.cell.final, it.reason)
        if not passes(report):
            print("STOP: dry-run tidak PASS"); return 2
        if not a.apply:
            print("DRY-RUN PASS: tidak ada write."); return 0
        ok, checks, n = execute(conn, cells, only)
        print("VALIDATION", json.dumps(checks, indent=1, ensure_ascii=False))
        if ok:
            conn.commit(); print(f"COMMIT: {n}"); return 0
        conn.rollback(); print("ROLLBACK"); return 3
    finally:
        conn.close()


if __name__ == "__main__":
    sys.exit(main())
